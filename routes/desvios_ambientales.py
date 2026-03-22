from flask import Blueprint, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from extensions import mysql
from utils.helpers import save_image, delete_image_file, sp_exec, sp_one, admin_required, modulo_required, get_notif_count
import io, os
from datetime import datetime

da_bp = Blueprint('da', __name__)

def get_maestros():
    cur = mysql.connection.cursor()
    cur.execute("SELECT * FROM tbl_areareportante ORDER BY areareportante")
    areas_rep = cur.fetchall()
    cur.execute("SELECT * FROM tbl_arearesponsable ORDER BY arearesponsable")
    areas_res = cur.fetchall()
    cur.execute("SELECT * FROM tbl_ubicacion ORDER BY ubicacion")
    ubicaciones = cur.fetchall()
    cur.execute("SELECT * FROM tbl_riesgo")
    riesgos = cur.fetchall()
    cur.execute("SELECT * FROM tbl_descripciontipo ORDER BY descripciontipo")
    tipos = cur.fetchall()
    cur.execute("SELECT * FROM tbl_estado ORDER BY orden")
    estados = cur.fetchall()
    cur.close()
    return areas_rep, areas_res, ubicaciones, riesgos, tipos, estados

# ── Dashboard ────────────────────────────────────────────────────────────────

@da_bp.route('/dashboard')
@admin_required
@modulo_required('DASHBOARD')
def dashboard():
    cur = mysql.connection.cursor()
    stats = sp_one(cur, 'sp_dashboardstats')
    cur.close()
    cur = mysql.connection.cursor()
    notifs = sp_exec(cur, 'sp_notificaciones', (session['usuario_rol'],))
    cur.close()
    cur = mysql.connection.cursor()
    todos = sp_exec(cur, 'sp_listarregistros', (None,))
    cur.close()
    registros_pendientes = [r for r in todos if r.get('estado') == 'Pendiente']
    total_pendientes = len(registros_pendientes)
    return render_template('desvios_ambientales/dashboard.html',
        stats=stats or {'total':0,'culminados':0,'en_proceso':0,'pendientes':0},
        notifs=notifs, notif_count=get_notif_count(),
        registros=registros_pendientes, total=total_pendientes,
        page=1, total_pages=1, per_page=total_pendientes or 10,
        estado_filter='Pendiente', personal_filter='')

# ── Registrar / Listar ───────────────────────────────────────────────────────

@da_bp.route('/registrar')
@admin_required
@modulo_required('DESVIOS')
def registrar():
    estado_filter  = request.args.get('estado', '')
    personal_filter = request.args.get('personal', '')
    page    = request.args.get('page', 1, type=int)
    per_page = 10

    cur = mysql.connection.cursor()
    todos_registros = sp_exec(cur, 'sp_listarregistros', (None,))
    cur.close()
    estados_unicos = sorted(list(set([r.get('estado','') for r in todos_registros if r.get('estado')])))

    cur = mysql.connection.cursor()
    registros = sp_exec(cur, 'sp_listarregistros', (estado_filter or None,))
    cur.close()

    if personal_filter:
        registros = [r for r in registros if r.get('personalresponsable','').lower() == personal_filter.lower()]

    orden_estados = {'Pendiente':1,'En Proceso':2,'Enviado':3,'En Revision':4,'Culminado':5,'Rechazado':6,'Cerrado':7}
    registros_ordenados = sorted(registros, key=lambda x: orden_estados.get(x.get('estado',''), 999))

    total       = len(registros_ordenados)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page        = max(1, min(page, total_pages))
    start       = (page - 1) * per_page
    registros_pagina = registros_ordenados[start:start + per_page]

    areas_rep, areas_res, ubicaciones, riesgos, tipos, estados = get_maestros()
    return render_template('desvios_ambientales/desvios.html',
        registros=registros_pagina,
        areas_rep=areas_rep, areas_res=areas_res,
        ubicaciones=ubicaciones, riesgos=riesgos, tipos=tipos, estados=estados,
        estado_filter=estado_filter, personal_filter=personal_filter,
        notif_count=get_notif_count(),
        page=page, total_pages=total_pages, total=total, per_page=per_page,
        estados_unicos=estados_unicos)

# ── Estadísticas ─────────────────────────────────────────────────────────────

@da_bp.route('/estadisticas')
@admin_required
@modulo_required('ESTADISTICAS')
def estadisticas():
    return render_template('desvios_ambientales/estadisticas.html', notif_count=get_notif_count())

@da_bp.route('/estadisticas/areas')
@admin_required
@modulo_required('ESTADISTICAS')
def estadisticas_areas():
    fecha_ini = request.args.get('fecha_ini') or None
    fecha_fin = request.args.get('fecha_fin') or None
    cur = mysql.connection.cursor()
    stats = sp_exec(cur, 'sp_estadisticasareas', (fecha_ini, fecha_fin))
    cur.close()
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify([dict(r) for r in stats])
    return redirect(url_for('da.estadisticas'))

@da_bp.route('/estadisticas/ccta')
@admin_required
@modulo_required('ESTADISTICAS')
def estadisticas_ccta():
    fecha_ini = request.args.get('fecha_ini') or None
    fecha_fin = request.args.get('fecha_fin') or None
    cur = mysql.connection.cursor()
    stats = sp_exec(cur, 'sp_estadisticasccta', (fecha_ini, fecha_fin))
    cur.close()
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify([dict(r) for r in stats])
    return redirect(url_for('da.estadisticas'))

@da_bp.route('/estadisticas/tipos')
@admin_required
@modulo_required('ESTADISTICAS')
def estadisticas_tipos():
    fecha_ini = request.args.get('fecha_ini') or None
    fecha_fin = request.args.get('fecha_fin') or None
    cur = mysql.connection.cursor()
    stats = sp_exec(cur, 'sp_estadisticas_tipos_pendientes', (fecha_ini, fecha_fin))
    cur.close()
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify([dict(r) for r in stats])
    return redirect(url_for('da.estadisticas'))

# ── CRUD Registros ───────────────────────────────────────────────────────────

@da_bp.route('/desvios')
@admin_required
def desvios_redirect():
    return redirect(url_for('da.registrar', **request.args))

@da_bp.route('/desvios/crear', methods=['POST'])
@admin_required
@modulo_required('DESVIOS')
def crear_registro():
    try:
        cur = mysql.connection.cursor()
        corr = sp_one(cur, 'sp_siguientecorrelativo')
        cur.close()
        codigo = corr['correlativo'] if corr else '001-01'

        cur = mysql.connection.cursor()
        result = sp_one(cur, 'sp_crearregistro', (
            codigo,
            request.form.get('fecha_inicio') or datetime.now().strftime('%Y-%m-%d'),
            request.form.get('fecha_ejecucion') or None,
            request.form['descripcion'], request.form.get('accion',''),
            int(request.form['area_reportante']), int(request.form['area_responsable']),
            int(request.form['ubicacion']), int(request.form['riesgo']), int(request.form['tipo']),
            1, session['usuario_rol'],
            request.form.get('personal_responsable',''),
            int(request.form['ccta_responsable']) if request.form.get('ccta_responsable') else 0,
            request.form.get('dni_responsable','').strip()
        ))
        mysql.connection.commit()
        cur.close()
        rid = result['idregistro'] if result else None

        for f in request.files.getlist('evidencias')[:5]:
            if f and f.filename:
                ruta, nombre, kb = save_image(f, 'static/uploads', 'evidencias')
                if ruta:
                    cur = mysql.connection.cursor()
                    sp_exec(cur, 'sp_guardarimagen', (rid, session['usuario_rol'], 1, 1, ruta, nombre, kb))
                    mysql.connection.commit()
                    cur.close()

        for f in request.files.getlist('levantamientos')[:5]:
            if f and f.filename:
                ruta, nombre, kb = save_image(f, 'static/uploads', 'levantamientos')
                if ruta:
                    cur = mysql.connection.cursor()
                    sp_exec(cur, 'sp_guardarimagen', (rid, session['usuario_rol'], 2, 1, ruta, nombre, kb))
                    mysql.connection.commit()
                    cur.close()

        cur = mysql.connection.cursor()
        cur.execute("SELECT ur.idusuariorol FROM tbl_usuariorol ur JOIN tbl_roles r ON r.idroles=ur.idroles WHERE r.nombrerol IN ('Supervisor','Trabajador')")
        sups = cur.fetchall()
        cur.close()
        for s in sups:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'sp_crearnotificacion', (s['idusuariorol'], f'Nuevo reporte {codigo} creado', 'info', rid))
            mysql.connection.commit()
            cur.close()

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': 'Reporte creado exitosamente', 'registro_id': rid, 'codigo': codigo})
        flash('Reporte creado exitosamente', 'success')
    except Exception as e:
        import traceback; traceback.print_exc()
        msg = f'Error al crear reporte: {str(e)}'
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': msg}), 400
        flash(msg, 'error')
    return redirect(url_for('da.registrar'))

@da_bp.route('/desvios/editar/<rid>', methods=['POST'])
@admin_required
@modulo_required('DESVIOS')
def editar_registro(rid):
    try:
        cur = mysql.connection.cursor()
        sp_exec(cur, 'sp_actualizarregistro', (
            rid,
            request.form.get('fecha_inicio') or datetime.now().strftime('%Y-%m-%d'),
            request.form.get('fecha_ejecucion') or None,
            request.form['descripcion'], request.form.get('accion',''),
            int(request.form['area_reportante']), int(request.form['area_responsable']),
            int(request.form['ubicacion']), int(request.form['riesgo']), int(request.form['tipo']),
            int(request.form['estado']),
            request.form.get('personal_responsable',''),
            int(request.form['ccta_responsable']) if request.form.get('ccta_responsable','').strip() not in ('','0','None') else 0,
            request.form.get('dni_responsable','').strip()
        ))
        mysql.connection.commit()
        cur.close()

        imagenes_eliminar = request.form.get('imagenes_eliminar','')
        if imagenes_eliminar:
            for imagen_id in [i.strip() for i in imagenes_eliminar.split(',') if i.strip()]:
                cur = mysql.connection.cursor()
                cur.execute("SELECT rutaimagen FROM tbl_imagenregistro WHERE idimagen=%s", (imagen_id,))
                img_row = cur.fetchone()
                cur.close()
                cur = mysql.connection.cursor()
                sp_exec(cur, 'sp_eliminarimagen', (imagen_id,))
                mysql.connection.commit()
                cur.close()
                if img_row:
                    delete_image_file(img_row.get('rutaimagen',''))

        for f in request.files.getlist('nuevas_evidencias')[:5]:
            if f and f.filename:
                ruta, nombre, kb = save_image(f, 'static/uploads', 'evidencias')
                if ruta:
                    cur = mysql.connection.cursor()
                    sp_exec(cur, 'sp_guardarimagen', (rid, session['usuario_rol'], 1, 1, ruta, nombre, kb))
                    mysql.connection.commit()
                    cur.close()

        for f in request.files.getlist('nuevos_levantamientos')[:5]:
            if f and f.filename:
                ruta, nombre, kb = save_image(f, 'static/uploads', 'levantamientos')
                if ruta:
                    cur = mysql.connection.cursor()
                    sp_exec(cur, 'sp_guardarimagen', (rid, session['usuario_rol'], 2, 1, ruta, nombre, kb))
                    mysql.connection.commit()
                    cur.close()

        flash('Reporte actualizado', 'success')
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': 'Reporte actualizado exitosamente'})
    except Exception as e:
        import traceback; traceback.print_exc()
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': f'Error: {str(e)}', 'traceback': traceback.format_exc()}), 400
        flash(f'Error: {str(e)}', 'error')
    return redirect(url_for('da.registrar'))

@da_bp.route('/desvios/eliminar/<rid>', methods=['POST'])
@admin_required
@modulo_required('DESVIOS')
def eliminar_registro(rid):
    try:
        cur = mysql.connection.cursor()
        cur.execute("SELECT rutaimagen FROM tbl_imagenregistro WHERE idregistro=%s", (rid,))
        imagenes = cur.fetchall()
        cur.execute("DELETE FROM tbl_historialaprobacion WHERE idregistro=%s", (rid,))
        cur.execute("DELETE FROM tbl_notificacion WHERE idregistro=%s", (rid,))
        cur.execute("DELETE FROM tbl_registro WHERE idregistro=%s", (rid,))
        mysql.connection.commit()
        cur.close()
        base_dir = os.path.join(os.path.dirname(__file__), '..')
        for img in imagenes:
            ruta = img.get('rutaimagen','')
            if ruta:
                filepath = os.path.normpath(os.path.join(base_dir, 'static', ruta))
                try:
                    if os.path.exists(filepath): os.remove(filepath)
                except: pass
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True})
        flash('Registro eliminado', 'success')
    except Exception as e:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': str(e)}), 400
        flash(f'Error al eliminar: {str(e)}', 'error')
    return redirect(url_for('da.registrar'))

@da_bp.route('/desvios/detalle/<rid>')
@admin_required
@modulo_required('DESVIOS')
def detalle_registro(rid):
    cur = mysql.connection.cursor()
    registro = sp_one(cur, 'sp_detalleregistro', (rid,))
    cur.close()
    cur = mysql.connection.cursor()
    imagenes = sp_exec(cur, 'sp_imagenesregistro', (rid,))
    cur.close()

    def serialize(obj):
        if obj is None: return {}
        out = {}
        for k, v in obj.items():
            key = k.lower()
            out[key] = v.strftime('%Y-%m-%d') if hasattr(v, 'strftime') else (v if v is not None else '')
        return out

    imgs_serial = [serialize(i) for i in imagenes if i.get('idEstadoImagen') != 3]
    evidencias     = [i for i in imgs_serial if i.get('idtipoimagen') == 1]
    levantamientos = [i for i in imgs_serial if i.get('idtipoimagen') == 2]
    return jsonify({'registro': serialize(registro), 'evidencias': evidencias, 'levantamientos': levantamientos})

@da_bp.route('/desvios/validar/<rid>', methods=['POST'])
@admin_required
@modulo_required('DESVIOS')
def validar_levantamiento(rid):
    try:
        decision   = request.form.get('decision')
        comentario = request.form.get('comentario','')

        if decision == 'APROBADA':
            imagenes_ids = request.form.get('imagenes_ids','')
            if not imagenes_ids:
                flash('No se recibieron imagenes para aprobar', 'error')
                return redirect(url_for('da.registrar'))
            ids_list = [i.strip() for i in imagenes_ids.split(',') if i.strip()]
            imagenes_ids_rechazar = request.form.get('imagenes_ids_rechazar','')
            if imagenes_ids_rechazar:
                ids_rechazar_list = [i.strip() for i in imagenes_ids_rechazar.split(',') if i.strip()]
                for imagen_id in [i for i in ids_rechazar_list if i not in ids_list]:
                    cur = mysql.connection.cursor()
                    cur.execute("SELECT rutaimagen FROM tbl_imagenregistro WHERE idimagen=%s", (imagen_id,))
                    img_row = cur.fetchone()
                    cur.close()
                    cur = mysql.connection.cursor()
                    sp_exec(cur, 'sp_eliminarimagen', (imagen_id,))
                    mysql.connection.commit()
                    cur.close()
                    if img_row:
                        delete_image_file(img_row.get('rutaimagen',''))
            for imagen_id in ids_list:
                cur = mysql.connection.cursor()
                sp_exec(cur, 'sp_validarimagen', (imagen_id, rid, session['usuario_rol'], 'APROBADA', comentario))
                mysql.connection.commit()
                cur.close()
            flash(f'{len(ids_list)} imagen(es) aprobada(s) correctamente', 'success')
        else:
            imagenes_ids_rechazar = request.form.get('imagenes_ids_rechazar','')
            if not imagenes_ids_rechazar:
                flash('No se recibieron imagenes para rechazar', 'error')
                return redirect(url_for('da.registrar'))
            ids_list = [i.strip() for i in imagenes_ids_rechazar.split(',') if i.strip()]
            # Obtener rutas antes de marcar como rechazadas
            rutas_rechazadas = {}
            if ids_list:
                cur = mysql.connection.cursor()
                cur.execute("SELECT idimagen, rutaimagen FROM tbl_imagenregistro WHERE idimagen IN (%s)" % ','.join(['%s']*len(ids_list)), ids_list)
                for row in cur.fetchall():
                    rutas_rechazadas[str(row['idimagen'])] = row.get('rutaimagen','')
                cur.close()
            for imagen_id in ids_list:
                cur = mysql.connection.cursor()
                sp_exec(cur, 'sp_validarimagen', (imagen_id, rid, session['usuario_rol'], 'RECHAZADA', comentario))
                mysql.connection.commit()
                cur.close()
                delete_image_file(rutas_rechazadas.get(imagen_id,''))
            flash(f'{len(ids_list)} imagen(es) rechazada(s). El supervisor debe volver a subir imagenes', 'warning')

        imagenes_ids_rechazar = request.form.get('imagenes_ids_rechazar','')
        if imagenes_ids_rechazar:
            ids_notif = [i.strip() for i in imagenes_ids_rechazar.split(',') if i.strip()]
            cur = mysql.connection.cursor()
            cur.execute("SELECT DISTINCT idusuariorol FROM tbl_imagenregistro WHERE idimagen IN (%s)" % ','.join(['%s']*len(ids_notif)), ids_notif)
            usuarios = cur.fetchall()
            cur.close()
            for usuario in usuarios:
                msg = f'Tu conjunto de imagenes fue {"APROBADO" if decision=="APROBADA" else "RECHAZADO"}'
                if comentario: msg += f': {comentario}'
                cur = mysql.connection.cursor()
                sp_exec(cur, 'sp_crearnotificacion', (usuario['idusuariorol'], msg, 'success' if decision=='APROBADA' else 'warning', rid))
                mysql.connection.commit()
                cur.close()

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': f'Imagenes {"aprobadas" if decision=="APROBADA" else "rechazadas"} exitosamente'})
    except Exception as e:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': f'Error: {str(e)}'}), 400
        flash(f'Error: {str(e)}', 'error')
    return redirect(url_for('da.registrar'))

@da_bp.route('/exportar')
@admin_required
@modulo_required('DESVIOS')
def exportar_excel():
    try:
        import openpyxl, tempfile
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage

        cur = mysql.connection.cursor()
        rows = sp_exec(cur, 'sp_exportarregistros')
        cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Desvios Ambientales"
        hfill = PatternFill("solid", fgColor="1a7a3c")
        hfont = Font(bold=True, color="FFFFFF", size=11)
        thin  = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))

        headers = ['Codigo','Fecha','Fecha Ejecucion','Area Reportante','Ubicacion',
                   'Descripcion','Tipo','Riesgo','Estado','Accion',
                   'Area Responsable','Personal Responsable','Fecha Creacion',
                   'Evidencias','Levantamientos']
        col_w = [12,14,16,20,30,40,30,12,14,40,22,22,18,25,25]

        for i,(h,w) in enumerate(zip(headers, col_w), 1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill = hfill; c.font = hfont; c.border = thin
            c.alignment = Alignment(horizontal='center', vertical='center')
            ws.column_dimensions[get_column_letter(i)].width = w

        alt = PatternFill("solid", fgColor="f0f9f4")
        current_row = 2
        temp_files = []

        for row in rows:
            cur = mysql.connection.cursor()
            imagenes = sp_exec(cur, 'sp_imagenesregistro', (row.get('IdRegistro',''),))
            cur.close()
            evidencias = []; levantamientos = []
            for img in imagenes:
                if img.get('idEstadoImagen') != 3:
                    if img.get('idTipoImagen') == 1: evidencias.append(img.get('RutaImagen',''))
                    elif img.get('idTipoImagen') == 2: levantamientos.append(img.get('RutaImagen',''))

            max_images = max(len(evidencias), len(levantamientos), 1)
            vals = [row.get('Codigo',''), row.get('Fecha',''), row.get('FechaEjecucion',''),
                    row.get('AreaReportante',''), row.get('Ubicacion',''), row.get('Descripcion',''),
                    row.get('DescripcionTipo',''), row.get('Riesgo',''), row.get('Estado',''),
                    row.get('Accion',''), row.get('AreaResponsable',''),
                    row.get('PersonalResponsable',''), row.get('FechaCreacion','')]

            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=current_row, column=ci, value=str(v) if v else '')
                c.border = thin
                c.alignment = Alignment(vertical='center', wrap_text=True)
                if current_row % 2 == 0: c.fill = alt
                if max_images > 1:
                    ws.merge_cells(start_row=current_row, start_column=ci,
                                   end_row=current_row+max_images-1, end_column=ci)

            for img_idx in range(max_images):
                img_row = current_row + img_idx
                ws.row_dimensions[img_row].height = 100
                for col_idx, img_list in [(14, evidencias), (15, levantamientos)]:
                    if img_idx < len(img_list):
                        img_path = img_list[img_idx]
                        possible = [img_path,
                                    os.path.join('ecosupervisor', img_path),
                                    os.path.join('static', img_path.replace('static/','')),
                                    os.path.join('ecosupervisor','static', img_path.replace('static/',''))]
                        full_path = next((p for p in possible if os.path.exists(p)), None)
                        if full_path:
                            try:
                                pil_img = PILImage.open(full_path)
                                pil_img.thumbnail((80,80), PILImage.Resampling.LANCZOS)
                                tf = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
                                temp_files.append(tf.name); tf.close()
                                pil_img.save(tf.name, 'PNG')
                                xl_img = XLImage(tf.name)
                                ws.add_image(xl_img, f'{get_column_letter(col_idx)}{img_row}')
                            except Exception as e:
                                ws.cell(row=img_row, column=col_idx, value=f'Error: {str(e)[:30]}')
                        else:
                            ws.cell(row=img_row, column=col_idx, value='No encontrada')
                    c = ws.cell(row=img_row, column=col_idx)
                    c.border = thin
                    if current_row % 2 == 0: c.fill = alt
            current_row += max_images

        out = io.BytesIO()
        wb.save(out); out.seek(0)
        for tp in temp_files:
            try:
                if os.path.exists(tp): os.remove(tp)
            except: pass
        fname = f"desvios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('da.registrar'))
