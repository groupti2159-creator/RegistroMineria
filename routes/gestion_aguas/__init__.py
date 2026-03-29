from flask import Blueprint, redirect, url_for, session
from functools import wraps
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

gestion_aguas_bp = Blueprint('gestion_aguas', __name__)

def _login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated

def _num(val, default=None):
    try:
        return float(val) if val not in (None, '', 'None') else default
    except (ValueError, TypeError):
        return default

def _serializar(row):
    out = {}
    for k, v in row.items():
        if hasattr(v, 'strftime'):
            out[k.lower()] = v.strftime('%Y-%m-%d')
        elif isinstance(v, (int, float)):
            out[k.lower()] = v
        elif v is None:
            out[k.lower()] = None
        else:
            out[k.lower()] = str(v)
    return out

# ── Helpers Excel compartidos ────────────────────────────
def excel_header_style():
    return (
        PatternFill("solid", fgColor="1a7a3c"),
        Font(bold=True, color="FFFFFF", size=11),
        Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'),  bottom=Side(style='thin')
        )
    )

def excel_escribir_headers(ws, headers, col_widths):
    fill, font, border = excel_header_style()
    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = font
        c.border = border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w

def excel_escribir_fila(ws, row_num, valores):
    alt    = PatternFill("solid", fgColor="f0f9f4")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'),  bottom=Side(style='thin'))
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=row_num, column=i, value=v if v is not None else '')
        c.border = border
        c.alignment = Alignment(vertical='center')
        if row_num % 2 == 0:
            c.fill = alt

# Importar rutas — debe ir al final
from routes.gestion_aguas import dashboard, monitoreo_ambiental, reporte_ana, exportar, exportar_ana