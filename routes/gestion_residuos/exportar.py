from flask import request, jsonify, send_file
from extensions import mysql
from utils.helpers import admin_required
from decimal import Decimal
from datetime import date, datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from routes.gestion_residuos import gr_bp


# ── Helpers Excel ─────────────────────────────────────────────────────────────

def _serialize(v):
    if isinstance(v, Decimal): return float(v)
    if isinstance(v, (date, datetime)): return v.isoformat()
    return v


def _xl_styles():
    hfill   = PatternFill("solid", fgColor="1a7a3c")
    hfont   = Font(bold=True, color="FFFFFF", size=11)
    altfill = PatternFill("solid", fgColor="f0f9f4")
    thin    = Border(left=Side(style='thin'), right=Side(style='thin'),
                     top=Side(style='thin'),  bottom=Side(style='thin'))
    return hfill, hfont, altfill, thin


def _xl_header(ws, headers, widths):
    hfill, hfont, _, thin = _xl_styles()
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = hfill; c.font = hfont; c.border = thin
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20


def _xl_row(ws, row_num, values, thin, altfill):
    for ci, v in enumerate(values, 1):
        c = ws.cell(row=row_num, column=ci, value=v)
        c.border = thin
        c.alignment = Alignment(vertical='center', wrap_text=True)
        if row_num % 2 == 0:
            c.fill = altfill


def _fecha_str(v):
    if not v: return ''
    s = str(v)[:10]
    parts = s.split('-')
    return f"{parts[2]}/{parts[1]}/{parts[0]}" if len(parts) == 3 else s


_GEN_SECTIONS = [
    {"id": "gen",  "label": "RRSS Generales",       "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "org",  "label": "RRSS Orgánicos",        "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "ind",  "label": "RRSS Industriales",     "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "hosp", "label": "Hospitalarios",         "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "ace",  "label": "Aceites",               "unit": "gal", "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "bor",  "label": "Borras",                "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "cop",  "label": "Copelas",               "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "tier", "label": "Tierra contaminada",    "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "var",  "label": "Varios",                "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "tar",  "label": "Tarros y res. menores", "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "chat", "label": "Chatarra mayor",        "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "plas", "label": "RRSS Plástico",         "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "vid",  "label": "RRSS Vidrio",           "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
    {"id": "pap",  "label": "RRSS Papel y cartón",   "unit": "Kg",
     "col_keys": ["pap_sunec","pap_coripuno","pap_antioquia","papa_smaria"],
     "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","SANTA MARÍA"]},
    {"id": "mad",  "label": "RRSS Madera",           "unit": "Kg",  "plants": ["SUNEC","CORI PUNO","ANTIOQUIA","PARCOY"]},
]


def _gen_col_keys(sec):
    if 'col_keys' in sec:
        return sec['col_keys']
    sid = sec['id']
    return [f"{sid}_sunec", f"{sid}_coripuno", f"{sid}_antioquia", f"{sid}_parcoy"]


# ── Rutas de exportación ──────────────────────────────────────────────────────

@gr_bp.route('/residuos/comercializable/exportar')
@admin_required
def comercializable_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        cur = mysql.connection.cursor()
        sql = "SELECT * FROM tbl_comercializable"
        conds, params = [], []
        if desde: conds.append("fecha >= %s"); params.append(desde)
        if hasta: conds.append("fecha <= %s"); params.append(hasta)
        if conds: sql += " WHERE " + " AND ".join(conds)
        sql += " ORDER BY fecha DESC"
        cur.execute(sql, params)
        rows = cur.fetchall() or []
        cur.close()

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Comercializable"
        headers = ['Fecha','Supervisor','Tipo de Residuo','Precio (S/)','N° Guía','Factura','Peso (Kg)','Observación','Creado por','Fecha Registro']
        widths  = [14, 25, 25, 14, 16, 16, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()
        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [_fecha_str(r.get('fecha')), r.get('supervisor',''), r.get('tipo_residuo',''),
                float(r.get('precio') or 0), r.get('nro_guia',''), r.get('factura',''),
                float(r.get('peso') or 0), r.get('observacion',''), r.get('creado_por',''),
                _fecha_str(r.get('fecha_registro'))], thin, altfill)

        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out, download_name=f"comercializable_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@gr_bp.route('/residuos/matpel/exportar')
@admin_required
def matpel_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        cur = mysql.connection.cursor()
        sql = "SELECT * FROM tbl_matpel"
        conds, params = [], []
        if desde: conds.append("fecha >= %s"); params.append(desde)
        if hasta: conds.append("fecha <= %s"); params.append(hasta)
        if conds: sql += " WHERE " + " AND ".join(conds)
        sql += " ORDER BY fecha DESC"
        cur.execute(sql, params)
        rows = cur.fetchall() or []
        cur.close()

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "MATPEL"
        headers = ['Fecha','Supervisor','Tipo de Residuo','Costo Viaje (S/)','Precio (S/)','N° Guía','Factura','Volumen (m³)','Peso (Kg)','Observación','Creado por','Fecha Registro']
        widths  = [14, 25, 25, 16, 14, 16, 16, 14, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()
        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [_fecha_str(r.get('fecha')), r.get('supervisor',''), r.get('tipo_residuo',''),
                float(r.get('costo_viaje') or 0), float(r.get('precio') or 0),
                r.get('nro_guia',''), r.get('factura',''),
                float(r.get('volumen') or 0), float(r.get('peso') or 0),
                r.get('observacion',''), r.get('creado_por',''),
                _fecha_str(r.get('fecha_registro'))], thin, altfill)

        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out, download_name=f"matpel_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@gr_bp.route('/residuos/compostaje/exportar')
@admin_required
def compostaje_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        cur = mysql.connection.cursor()
        sql = "SELECT * FROM tbl_compostaje"
        conds, params = [], []
        if desde: conds.append("fecha >= %s"); params.append(desde)
        if hasta: conds.append("fecha <= %s"); params.append(hasta)
        if conds: sql += " WHERE " + " AND ".join(conds)
        sql += " ORDER BY fecha DESC"
        cur.execute(sql, params)
        rows = cur.fetchall() or []
        cur.close()

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Compostaje"
        headers = ['Fecha','Preparación (Kg)','Cosecha (Kg)','Costo Viaje (S/)','Precio (S/)','N° Guía','Factura','Volumen (m³)','Peso (Kg)','Observación','Creado por','Fecha Registro']
        widths  = [14, 18, 16, 16, 14, 16, 16, 14, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()
        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [_fecha_str(r.get('fecha')),
                float(r.get('preparacion') or 0), float(r.get('cosecha') or 0),
                float(r.get('costo_viaje') or 0), float(r.get('precio') or 0),
                r.get('nro_guia',''), r.get('factura',''),
                float(r.get('volumen') or 0), float(r.get('peso') or 0),
                r.get('observacion',''), r.get('creado_por',''),
                _fecha_str(r.get('fecha_registro'))], thin, altfill)

        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out, download_name=f"compostaje_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@gr_bp.route('/residuos/generacion/exportar')
@admin_required
def generacion_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        cur = mysql.connection.cursor()
        sql = "SELECT * FROM tbl_generacion_diaria"
        conds, params = [], []
        if desde: conds.append("fecha >= %s"); params.append(desde)
        if hasta: conds.append("fecha <= %s"); params.append(hasta)
        if conds: sql += " WHERE " + " AND ".join(conds)
        sql += " ORDER BY fecha DESC"
        cur.execute(sql, params)
        rows = cur.fetchall() or []
        cur.close()

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Generación Diaria"
        headers = ['Fecha', 'Gen. Diaria (t)']
        widths  = [14, 16]
        for sec in _GEN_SECTIONS:
            for plant in sec['plants']:
                headers.append(f"{sec['label']} - {plant} ({sec['unit']})")
                widths.append(28)
            headers.append(f"{sec['label']} - TOTAL ({sec['unit']})")
            widths.append(24)

        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()

        for i, r in enumerate(rows, 2):
            rd = {k.lower(): _serialize(v) for k, v in r.items()} if isinstance(r, dict) else r
            values = [_fecha_str(rd.get('fecha')), float(rd.get('generacion_t') or 0)]
            for sec in _GEN_SECTIONS:
                for ck in _gen_col_keys(sec):
                    values.append(float(rd.get(ck) or 0))
                values.append(float(rd.get(sec['id'] + '_total') or 0))
            _xl_row(ws, i, values, thin, altfill)

        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out, download_name=f"generacion_diaria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500
