from flask import request, redirect, url_for, flash, send_file
from extensions import mysql
from utils.helpers import sp_exec, admin_required, modulo_required
from datetime import datetime
import io, os
from routes.desvios_ambientales import da_bp


@da_bp.route('/exportar')
@admin_required
@modulo_required('DESVIOS')
def exportar_excel():
    try:
        import openpyxl, tempfile
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage

        cur = mysql.connection.cursor()
        rows = sp_exec(cur, 'sp_exportarregistros')
        cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Desvios Ambientales"
        hfill = PatternFill("solid", fgColor="1a7a3c")
        hfont = Font(bold=True, color="FFFFFF", size=11)
        thin  = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))

        headers = ['Codigo', 'Fecha', 'Fecha Ejecucion', 'Area Reportante', 'Ubicacion',
                   'Descripcion', 'Tipo', 'Riesgo', 'Estado', 'Accion',
                   'Area Responsable', 'Personal Responsable', 'Fecha Creacion',
                   'Evidencias', 'Levantamientos']
        col_w = [12, 14, 16, 20, 30, 40, 30, 12, 14, 40, 22, 22, 18, 25, 25]

        for i, (h, w) in enumerate(zip(headers, col_w), 1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill = hfill; c.font = hfont; c.border = thin
            c.alignment = Alignment(horizontal='center', vertical='center')
            ws.column_dimensions[get_column_letter(i)].width = w

        alt = PatternFill("solid", fgColor="f0f9f4")
        current_row = 2
        temp_files = []

        for row in rows:
            cur = mysql.connection.cursor()
            imagenes = sp_exec(cur, 'sp_imagenesregistro', (row.get('IdRegistro', ''),))
            cur.close()
            evidencias = []; levantamientos = []
            for img in imagenes:
                if img.get('idEstadoImagen') != 3:
                    if img.get('idTipoImagen') == 1: evidencias.append(img.get('RutaImagen', ''))
                    elif img.get('idTipoImagen') == 2: levantamientos.append(img.get('RutaImagen', ''))

            max_images = max(len(evidencias), len(levantamientos), 1)
            vals = [row.get('Codigo', ''), row.get('Fecha', ''), row.get('FechaEjecucion', ''),
                    row.get('AreaReportante', ''), row.get('Ubicacion', ''), row.get('Descripcion', ''),
                    row.get('DescripcionTipo', ''), row.get('Riesgo', ''), row.get('Estado', ''),
                    row.get('Accion', ''), row.get('AreaResponsable', ''),
                    row.get('PersonalResponsable', ''), row.get('FechaCreacion', '')]

            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=current_row, column=ci, value=str(v) if v else '')
                c.border = thin
                c.alignment = Alignment(vertical='center', wrap_text=True)
                if current_row % 2 == 0: c.fill = alt
                if max_images > 1:
                    ws.merge_cells(start_row=current_row, start_column=ci,
                                   end_row=current_row + max_images - 1, end_column=ci)

            for img_idx in range(max_images):
                img_row = current_row + img_idx
                ws.row_dimensions[img_row].height = 100
                for col_idx, img_list in [(14, evidencias), (15, levantamientos)]:
                    if img_idx < len(img_list):
                        img_path = img_list[img_idx]
                        possible = [img_path,
                                    os.path.join('ecosupervisor', img_path),
                                    os.path.join('static', img_path.replace('static/', '')),
                                    os.path.join('ecosupervisor', 'static', img_path.replace('static/', ''))]
                        full_path = next((p for p in possible if os.path.exists(p)), None)
                        if full_path:
                            try:
                                pil_img = PILImage.open(full_path)
                                pil_img.thumbnail((80, 80), PILImage.Resampling.LANCZOS)
                                tf = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
                                temp_files.append(tf.name); tf.close()
                                pil_img.save(tf.name, 'PNG')
                                ws.add_image(XLImage(tf.name), f'{get_column_letter(col_idx)}{img_row}')
                            except Exception as e:
                                ws.cell(row=img_row, column=col_idx, value=f'Error: {str(e)[:30]}')
                        else:
                            ws.cell(row=img_row, column=col_idx, value='No encontrada')
                    ws.cell(row=img_row, column=col_idx).border = thin
                    if current_row % 2 == 0:
                        ws.cell(row=img_row, column=col_idx).fill = alt
            current_row += max_images

        out = io.BytesIO()
        wb.save(out); out.seek(0)
        for tp in temp_files:
            try:
                if os.path.exists(tp): os.remove(tp)
            except Exception: pass

        fname = f"desvios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('da.registrar'))
