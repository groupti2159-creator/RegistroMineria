from flask import Blueprint, render_template, request, jsonify, session, send_file
from utils.helpers import admin_required, modulo_required, get_notif_count, sp_exec, sp_one, save_image, delete_image_file
import os
from extensions import mysql
from decimal import Decimal
from datetime import date, datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

gr_bp = Blueprint('gr', __name__)

# Secciones con su columna coripuno
SECTIONS_CP = [
    {"id": "gen",  "label": "RRSS Generales",
        "unit": "Kg",  "col": "gen_coripuno"},
    {"id": "org",  "label": "RRSS Orgánicos",
        "unit": "Kg",  "col": "org_coripuno"},
    {"id": "ind",  "label": "RRSS Industriales",
        "unit": "Kg",  "col": "ind_coripuno"},
    {"id": "hosp", "label": "Hospitalarios",
        "unit": "Kg",  "col": "hosp_coripuno"},
    {"id": "ace",  "label": "Aceites",
        "unit": "gal", "col": "ace_coripuno"},
    {"id": "bor",  "label": "Borras",
        "unit": "Kg",  "col": "bor_coripuno"},
    {"id": "cop",  "label": "Copelas",
        "unit": "Kg",  "col": "cop_coripuno"},
    {"id": "tier", "label": "Tierra contaminada",
        "unit": "Kg",  "col": "tier_coripuno"},
    {"id": "var",  "label": "Varios",
        "unit": "Kg",  "col": "var_coripuno"},
    {"id": "tar",  "label": "Tarros y res. menores",
        "unit": "Kg",  "col": "tar_coripuno"},
    {"id": "chat", "label": "Chatarra mayor",
        "unit": "Kg",  "col": "chat_coripuno"},
    {"id": "plas", "label": "RRSS Plástico",
        "unit": "Kg",  "col": "plas_coripuno"},
    {"id": "vid",  "label": "RRSS Vidrio",
        "unit": "Kg",  "col": "vid_coripuno"},
    {"id": "pap",  "label": "RRSS Papel y cartón",
        "unit": "Kg",  "col": "pap_coripuno"},
    {"id": "mad",  "label": "RRSS Madera",
        "unit": "Kg",  "col": "mad_coripuno"},
]


def _serialize(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


@gr_bp.route('/residuos/dashboard')
@admin_required
def dashboard():
    try:
        cur = mysql.connection.cursor()
        cols = ", ".join(["fecha"] + [s["col"] for s in SECTIONS_CP])
        cur.execute(
            f"SELECT {cols} FROM tbl_generacion_diaria ORDER BY fecha DESC")
        rows = cur.fetchall() or []
    finally:
        if 'cur' in locals() and cur:
            cur.close()

    # Convertir a lista de dicts serializables
    keys = ["fecha"] + [s["col"] for s in SECTIONS_CP]
    registros = []
    for row in rows:
        if isinstance(row, dict):
            r = {k.lower(): _serialize(v) for k, v in row.items()}
        else:
            r = {keys[i]: _serialize(row[i]) for i in range(len(keys))}
        # total kg CORI PUNO (sin aceites)
        r['total_cp'] = sum(
            float(r.get(s["col"], 0) or 0)
            for s in SECTIONS_CP if s["unit"] == "Kg"
        )
        registros.append(r)

    return render_template(
        'gestion_residuos/dashboard.html',
        notif_count=get_notif_count(),
        sections=SECTIONS_CP,
        registros=registros,
    )


@gr_bp.route('/residuos/comercializable')
@admin_required
@modulo_required('COMERCIALIZABLE')
def comercializable():
    try:
        cur = mysql.connection.cursor()
        cur.execute("SELECT * FROM tbl_comercializable ORDER BY fecha DESC")
        rows = cur.fetchall() or []
    finally:
        if 'cur' in locals() and cur:
            cur.close()
    registros = [{k.lower(): _serialize(v) for k, v in r.items()}
                 for r in rows]
    return render_template('gestion_residuos/comercializable.html',
                           notif_count=get_notif_count(), registros=registros)


@gr_bp.route('/residuos/comercializable/guardar', methods=['POST'])
@admin_required
def comercializable_guardar():
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            row = sp_one(cur, 'SP_Comercializable_Guardar', (
                d.get('fecha'), d.get('supervisor'), d.get('tipo'),
                float(d.get('precio')) if d.get(
                    'precio') else None, d.get('guia'), factura,
                float(d.get('peso')) if d.get('peso') else None, d.get(
                    'obs'), session.get('user_id')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True, 'id': row['id'] if row else None})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/comercializable/editar/<int:rid>', methods=['POST'])
@admin_required
def comercializable_editar(rid):
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        try:
            cur = mysql.connection.cursor()
            cur.execute(
                "SELECT factura FROM tbl_comercializable WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Comercializable_Editar', (
                rid, d.get('fecha'), d.get('supervisor'), d.get('tipo'),
                float(d.get('precio')) if d.get(
                    'precio') else None, d.get('guia'), factura,
                float(d.get('peso')) if d.get('peso') else None, d.get('obs')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/comercializable/eliminar/<int:rid>', methods=['POST'])
@admin_required
def comercializable_eliminar(rid):
    try:
        try:
            cur = mysql.connection.cursor()
            cur.execute(
                "SELECT factura FROM tbl_comercializable WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Comercializable_Eliminar', (rid,))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/matpel')
@admin_required
@modulo_required('DISPOSICION_MATPEL')
def matpel():
    try:
        cur = mysql.connection.cursor()
        cur.execute("SELECT * FROM tbl_matpel ORDER BY fecha DESC")
        rows = cur.fetchall() or []
    finally:
        if 'cur' in locals() and cur:
            cur.close()
    registros = [{k.lower(): _serialize(v) for k, v in r.items()}
                 for r in rows]
    return render_template('gestion_residuos/matpel.html',
                           notif_count=get_notif_count(), registros=registros)


@gr_bp.route('/residuos/matpel/guardar', methods=['POST'])
@admin_required
def matpel_guardar():
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            row = sp_one(cur, 'SP_Matpel_Guardar', (
                d.get('fecha'), d.get('supervisor'), d.get('tipo'),
                float(d.get('costo_viaje')) if d.get('costo_viaje') else None, float(
                    d.get('precio')) if d.get('precio') else None,
                d.get('guia'), factura,
                float(d.get('volumen')) if d.get('volumen') else None, float(
                    d.get('peso')) if d.get('peso') else None,
                d.get('obs'), session.get('user_id')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True, 'id': row['id'] if row else None})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/matpel/editar/<int:rid>', methods=['POST'])
@admin_required
def matpel_editar(rid):
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        try:
            cur = mysql.connection.cursor()
            cur.execute("SELECT factura FROM tbl_matpel WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Matpel_Editar', (
                rid, d.get('fecha'), d.get('supervisor'), d.get('tipo'),
                float(d.get('costo_viaje')) if d.get('costo_viaje') else None, float(
                    d.get('precio')) if d.get('precio') else None,
                d.get('guia'), factura,
                float(d.get('volumen')) if d.get('volumen') else None, float(
                    d.get('peso')) if d.get('peso') else None,
                d.get('obs')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/matpel/eliminar/<int:rid>', methods=['POST'])
@admin_required
def matpel_eliminar(rid):
    try:
        try:
            cur = mysql.connection.cursor()
            cur.execute("SELECT factura FROM tbl_matpel WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Matpel_Eliminar', (rid,))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/compostaje')
@admin_required
@modulo_required('COMPOSTAJE')
def compostaje():
    try:
        cur = mysql.connection.cursor()
        cur.execute("SELECT * FROM tbl_compostaje ORDER BY fecha DESC")
        rows = cur.fetchall() or []
    finally:
        if 'cur' in locals() and cur:
            cur.close()
    registros = [{k.lower(): _serialize(v) for k, v in r.items()}
                 for r in rows]
    return render_template('gestion_residuos/compostaje.html',
                           notif_count=get_notif_count(), registros=registros)


@gr_bp.route('/residuos/compostaje/guardar', methods=['POST'])
@admin_required
def compostaje_guardar():
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            row = sp_one(cur, 'SP_Compostaje_Guardar', (
                d.get('fecha'),
                float(d.get('preparacion')) if d.get('preparacion') else None, float(
                    d.get('cosecha')) if d.get('cosecha') else None,
                float(d.get('costo_viaje')) if d.get('costo_viaje') else None, float(
                    d.get('precio')) if d.get('precio') else None,
                d.get('guia'), factura,
                float(d.get('volumen')) if d.get('volumen') else None, float(
                    d.get('peso')) if d.get('peso') else None,
                d.get('obs'), session.get('user_id')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True, 'id': row['id'] if row else None})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/compostaje/editar/<int:rid>', methods=['POST'])
@admin_required
def compostaje_editar(rid):
    d = request.form
    factura = d.get('factura')
    if 'factura_file' in request.files and request.files['factura_file'].filename:
        try:
            cur = mysql.connection.cursor()
            cur.execute(
                "SELECT factura FROM tbl_compostaje WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        ruta, nombre, kb = save_image(
            request.files['factura_file'], 'static/uploads', 'facturas_residuos')
        if ruta:
            factura = ruta
    try:
        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Compostaje_Editar', (
                rid, d.get('fecha'),
                float(d.get('preparacion')) if d.get('preparacion') else None, float(
                    d.get('cosecha')) if d.get('cosecha') else None,
                float(d.get('costo_viaje')) if d.get('costo_viaje') else None, float(
                    d.get('precio')) if d.get('precio') else None,
                d.get('guia'), factura,
                float(d.get('volumen')) if d.get('volumen') else None, float(
                    d.get('peso')) if d.get('peso') else None,
                d.get('obs')
            ))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@gr_bp.route('/residuos/compostaje/eliminar/<int:rid>', methods=['POST'])
@admin_required
def compostaje_eliminar(rid):
    try:
        try:
            cur = mysql.connection.cursor()
            cur.execute(
                "SELECT factura FROM tbl_compostaje WHERE id = %s", (rid,))
            old = cur.fetchone()
            if old and isinstance(old, dict) and old.get('factura'):
                delete_image_file(old.get('factura'))
            elif old and isinstance(old, tuple) and old[0]:
                delete_image_file(old[0])
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        try:
            cur = mysql.connection.cursor()
            sp_exec(cur, 'SP_Compostaje_Eliminar', (rid,))
            mysql.connection.commit()
        finally:
            if 'cur' in locals() and cur:
                cur.close()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

# ── Helper estilos Excel ─────────────────────────────────────────────────────


def _xl_styles():
    hfill = PatternFill("solid", fgColor="1a7a3c")
    hfont = Font(bold=True, color="FFFFFF", size=11)
    altfill = PatternFill("solid", fgColor="f0f9f4")
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'),  bottom=Side(style='thin')
    )
    return hfill, hfont, altfill, thin


def _xl_header(ws, headers, widths):
    hfill, hfont, _, thin = _xl_styles()
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = hfill
        c.font = hfont
        c.border = thin
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20


def _xl_row(ws, row_num, values, thin, altfill):
    for ci, v in enumerate(values, 1):
        c = ws.cell(row=row_num, column=ci, value=v)
        c.border = thin
        c.alignment = Alignment(vertical='center', wrap_text=True)
        if row_num % 2 == 0:
            c.fill = altfill


def _fecha_str(v):
    if not v:
        return ''
    s = str(v)[:10]
    parts = s.split('-')
    return f"{parts[2]}/{parts[1]}/{parts[0]}" if len(parts) == 3 else s


# ── Export Comercializable ───────────────────────────────────────────────────
@gr_bp.route('/residuos/comercializable/exportar')
@admin_required
def comercializable_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        try:
            cur = mysql.connection.cursor()
            sql = "SELECT * FROM tbl_comercializable"
            conds, params = [], []
            if desde:
                conds.append("fecha >= %s")
                params.append(desde)
            if hasta:
                conds.append("fecha <= %s")
                params.append(hasta)
            if conds:
                sql += " WHERE " + " AND ".join(conds)
            sql += " ORDER BY fecha DESC"
            cur.execute(sql, params)
            rows = cur.fetchall() or []
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comercializable"
        headers = ['Fecha', 'Supervisor', 'Tipo de Residuo',
                   'Precio (S/)', 'N° Guía', 'Factura', 'Peso (Kg)', 'Observación', 'Creado por', 'Fecha Registro']
        widths = [14, 25, 25, 14, 16, 16, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()

        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [
                _fecha_str(r.get('fecha')),
                r.get('supervisor', ''), r.get('tipo_residuo', ''),
                float(r.get('precio') or 0), r.get('nro_guia', ''),
                r.get('factura', ''), float(r.get('peso') or 0),
                r.get('observacion', ''), r.get('creado_por', ''),
                _fecha_str(r.get('fecha_registro')),
            ], thin, altfill)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        fname = f"comercializable_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ── Export MATPEL ────────────────────────────────────────────────────────────
@gr_bp.route('/residuos/matpel/exportar')
@admin_required
def matpel_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        try:
            cur = mysql.connection.cursor()
            sql = "SELECT * FROM tbl_matpel"
            conds, params = [], []
            if desde:
                conds.append("fecha >= %s")
                params.append(desde)
            if hasta:
                conds.append("fecha <= %s")
                params.append(hasta)
            if conds:
                sql += " WHERE " + " AND ".join(conds)
            sql += " ORDER BY fecha DESC"
            cur.execute(sql, params)
            rows = cur.fetchall() or []
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MATPEL"
        headers = ['Fecha', 'Supervisor', 'Tipo de Residuo',
                   'Costo Viaje (S/)', 'Precio (S/)', 'N° Guía', 'Factura', 'Volumen (m³)', 'Peso (Kg)', 'Observación', 'Creado por', 'Fecha Registro']
        widths = [14, 25, 25, 16, 14, 16, 16, 14, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()

        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [
                _fecha_str(r.get('fecha')),
                r.get('supervisor', ''), r.get('tipo_residuo', ''),
                float(r.get('costo_viaje') or 0), float(r.get('precio') or 0),
                r.get('nro_guia', ''), r.get('factura', ''),
                float(r.get('volumen') or 0), float(r.get('peso') or 0),
                r.get('observacion', ''), r.get('creado_por', ''),
                _fecha_str(r.get('fecha_registro')),
            ], thin, altfill)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        fname = f"matpel_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ── Export Compostaje ────────────────────────────────────────────────────────
@gr_bp.route('/residuos/compostaje/exportar')
@admin_required
def compostaje_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        try:
            cur = mysql.connection.cursor()
            sql = "SELECT * FROM tbl_compostaje"
            conds, params = [], []
            if desde:
                conds.append("fecha >= %s")
                params.append(desde)
            if hasta:
                conds.append("fecha <= %s")
                params.append(hasta)
            if conds:
                sql += " WHERE " + " AND ".join(conds)
            sql += " ORDER BY fecha DESC"
            cur.execute(sql, params)
            rows = cur.fetchall() or []
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Compostaje"
        headers = ['Fecha', 'Preparación (Kg)', 'Cosecha (Kg)', 'Costo Viaje (S/)', 'Precio (S/)', 'N° Guía',
                   'Factura', 'Volumen (m³)', 'Peso (Kg)', 'Observación', 'Creado por', 'Fecha Registro']
        widths = [14, 18, 16, 16, 14, 16, 16, 14, 14, 35, 20, 18]
        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()

        for i, r in enumerate(rows, 2):
            _xl_row(ws, i, [
                _fecha_str(r.get('fecha')),
                float(r.get('preparacion') or 0), float(r.get('cosecha') or 0),
                float(r.get('costo_viaje') or 0), float(r.get('precio') or 0),
                r.get('nro_guia', ''), r.get('factura', ''),
                float(r.get('volumen') or 0), float(r.get('peso') or 0),
                r.get('observacion', ''), r.get('creado_por', ''),
                _fecha_str(r.get('fecha_registro')),
            ], thin, altfill)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        fname = f"compostaje_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ── Export Generación Diaria ─────────────────────────────────────────────────
# Secciones con sus 4 plantas y columna total
_GEN_SECTIONS = [
    {"id": "gen",  "label": "RRSS Generales",        "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "org",  "label": "RRSS Orgánicos",         "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "ind",  "label": "RRSS Industriales",      "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "hosp", "label": "Hospitalarios",          "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "ace",  "label": "Aceites",                "unit": "gal",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "bor",  "label": "Borras",                 "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "cop",  "label": "Copelas",                "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "tier", "label": "Tierra contaminada",     "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "var",  "label": "Varios",                 "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "tar",  "label": "Tarros y res. menores",  "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "chat", "label": "Chatarra mayor",         "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "plas", "label": "RRSS Plástico",          "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "vid",  "label": "RRSS Vidrio",            "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
    {"id": "pap",  "label": "RRSS Papel y cartón",    "unit": "Kg",
     "col_keys": ["pap_sunec", "pap_coripuno", "pap_antioquia", "papa_smaria"],
     "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "SANTA MARÍA"]},
    {"id": "mad",  "label": "RRSS Madera",            "unit": "Kg",
        "plants": ["SUNEC", "CORI PUNO", "ANTIOQUIA", "PARCOY"]},
]


def _gen_col_keys(sec):
    if 'col_keys' in sec:
        return sec['col_keys']
    sid = sec['id']
    return [f"{sid}_sunec", f"{sid}_coripuno", f"{sid}_antioquia", f"{sid}_parcoy"]


@gr_bp.route('/residuos/generacion/exportar')
@admin_required
def generacion_exportar():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    try:
        try:
            cur = mysql.connection.cursor()
            sql = "SELECT * FROM tbl_generacion_diaria"
            conds, params = [], []
            if desde:
                conds.append("fecha >= %s")
                params.append(desde)
            if hasta:
                conds.append("fecha <= %s")
                params.append(hasta)
            if conds:
                sql += " WHERE " + " AND ".join(conds)
            sql += " ORDER BY fecha DESC"
            cur.execute(sql, params)
            rows = cur.fetchall() or []
        finally:
            if 'cur' in locals() and cur:
                cur.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Generación Diaria"

        # Construir cabeceras: Fecha | Gen.Diaria(t) | [Sección - Planta x4 | Total] x15
        headers = ['Fecha', 'Gen. Diaria (t)']
        widths = [14, 16]
        for sec in _GEN_SECTIONS:
            for plant in sec['plants']:
                headers.append(f"{sec['label']} - {plant} ({sec['unit']})")
                widths.append(28)
            headers.append(f"{sec['label']} - TOTAL ({sec['unit']})")
            widths.append(24)

        _xl_header(ws, headers, widths)
        _, _, altfill, thin = _xl_styles()

        for i, r in enumerate(rows, 2):
            rd = {k.lower(): _serialize(v)
                  for k, v in r.items()} if isinstance(r, dict) else r
            values = [_fecha_str(rd.get('fecha')), float(
                rd.get('generacion_t') or 0)]
            for sec in _GEN_SECTIONS:
                col_keys = _gen_col_keys(sec)
                for ck in col_keys:
                    values.append(float(rd.get(ck) or 0))
                values.append(float(rd.get(sec['id'] + '_total') or 0))
            _xl_row(ws, i, values, thin, altfill)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        fname = f"generacion_diaria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(out, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    except Exception as e:
        return jsonify({'error': str(e)}), 500
