from flask import request, send_file
from extensions import mysql
from utils.helpers import sp_exec, get_notif_count
from routes.compromisos import compromisos_bp, _login_required
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date

def _estilo_header():
    return (
        PatternFill("solid", fgColor="1a7a3c"),
        Font(bold=True, color="FFFFFF", size=11),
        Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
    )

@compromisos_bp.route('/compromisos/exportar')
@_login_required
def exportar_compromisos():
    mes  = int(request.args.get('mes',  date.today().month))
    anio = int(request.args.get('anio', date.today().year))

    cur = mysql.connection.cursor()
    compromisos   = sp_exec(cur, 'sp_listarcompromisos')
    cumplimientos = sp_exec(cur, 'sp_obtenercumplimiento', (mes, anio))
    cur.close()

    cum_dict = {str(c.get('idcompromiso') or c.get('IDCOMPROMISO')): c
                for c in cumplimientos}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compromisos"

    fill, font, border = _estilo_header()
    alt  = PatternFill("solid", fgColor="f0f9f4")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                  top=Side(style='thin'),  bottom=Side(style='thin'))

    meses = {1:'Enero',2:'Febrero',3:'Marzo',4:'Abril',5:'Mayo',6:'Junio',
             7:'Julio',8:'Agosto',9:'Setiembre',10:'Octubre',
             11:'Noviembre',12:'Diciembre'}

    headers    = ['Permisos / Obligaciones Ambientales',
                  'Descripción / Justificación',
                  'Entidad Regulatoria',
                  'Tiene Evidencia',
                  'Supervisor',
                  'Observaciones',
                  f'Período']
    col_widths = [45, 45, 20, 16, 25, 35, 16]

    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = font
        c.border = border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    for i, comp in enumerate(compromisos, 2):
        id_str = str(comp.get('idcompromiso') or comp.get('IDCOMPROMISO', ''))
        cum    = cum_dict.get(id_str, {})

        tiene_ev   = cum.get('tiene_evidencia') or cum.get('TIENE_EVIDENCIA')
        supervisor = cum.get('idusuario') or cum.get('IDUSUARIO') or '—'
        obs        = cum.get('observaciones') or cum.get('OBSERVACIONES') or ''

        valores = [
            comp.get('nombre') or comp.get('NOMBRE', ''),
            comp.get('descripcion') or comp.get('DESCRIPCION', ''),
            comp.get('entidad_reguladora') or comp.get('ENTIDAD_REGULADORA', ''),
            'Sí' if tiene_ev else 'No',
            supervisor,
            obs,
            f"{meses.get(mes, mes)}/{anio}",
        ]

        for j, v in enumerate(valores, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = thin
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if i % 2 == 0:
                c.fill = alt

            # Colorear columna "Tiene Evidencia"
            if j == 4:
                if v == 'Sí':
                    c.font = Font(color="166534", bold=True)
                else:
                    c.font = Font(color="991b1b", bold=True)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    fname = f"compromisos_{meses.get(mes,'')}{anio}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(out, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True)